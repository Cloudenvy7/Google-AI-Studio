#!/usr/bin/env python3
"""Rebuild the Visual Second Brain relational substrate workbook from substrate_spec.json.

The spec (substrate_spec.json) is the committed source of truth for the schema and the seeded
foundation (RESOURCES + INTEGRATIONS). This script is deterministic: edit the spec, re-run, and
the .xlsx is regenerated identically.

Usage:
    python3 tools/build_substrate.py [--spec tools/substrate_spec.json] [--out Visual_Second_Brain_Substrate.xlsx]

Requires: openpyxl  (pip install openpyxl)

Why a script and not a committed .xlsx: the workbook is a binary artifact and cannot be diffed or
pushed through the text-only write path this repo uses. The spec is reviewable text; the workbook
is a generated projection of it — the same "sheets are truth, everything else is a projection"
discipline the system itself runs on.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    raise SystemExit("openpyxl is required. Install it with: pip install openpyxl")

TITLE_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FILL = PatternFill("solid", fgColor="374151")
TITLE_FONT = Font(bold=True, color="FFFFFF", size=13)
DESC_FONT = Font(italic=True, color="4B5563", size=10)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)


def _style_title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    c = ws.cell(row=1, column=1, value=text)
    c.fill, c.font = TITLE_FILL, TITLE_FONT
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22


def _style_desc(ws, text, ncols):
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncols, 1))
    c = ws.cell(row=2, column=1, value=text)
    c.font = DESC_FONT
    c.alignment = Alignment(vertical="center", wrap_text=True)


def _write_headers(ws, headers, row):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22


def build(spec_path: Path, out_path: Path) -> None:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)

    # 00_README
    ws = wb.create_sheet("00_README")
    _style_title(ws, spec.get("workbook_title", "Substrate"), 2)
    r = 3
    for left, right in spec.get("readme", []):
        ws.cell(row=r, column=1, value=left).font = Font(bold=True)
        ws.cell(row=r, column=2, value=right)
        r += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90

    seeds = spec.get("seeds", {})

    # Schema sheets
    for sheet in spec["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        headers = sheet["headers"]
        _style_title(ws, sheet["title"], len(headers))
        _style_desc(ws, sheet["description"], len(headers))
        _write_headers(ws, headers, row=3)
        rows = seeds.get(sheet.get("seed_key", ""), [])
        for ri, rowvals in enumerate(rows, start=4):
            for ci, val in enumerate(rowvals, start=1):
                ws.cell(row=ri, column=ci, value=val)

    # VOCABULARIES — one column per controlled list
    ws = wb.create_sheet("VOCABULARIES")
    vocab = spec.get("vocabularies", {})
    _style_title(ws, "CONTROLLED VOCABULARIES", max(len(vocab), 1))
    _style_desc(ws, "Canonical value lists for every enum column. Keep tight. Vocabulary explosion kills a relational system.", len(vocab))
    for ci, (col, values) in enumerate(vocab.items(), start=1):
        h = ws.cell(row=3, column=ci, value=col)
        h.fill, h.font = HEADER_FILL, HEADER_FONT
        for ri, v in enumerate(values, start=4):
            ws.cell(row=ri, column=ci, value=v)
        ws.column_dimensions[get_column_letter(ci)].width = 28
    ws.freeze_panes = "A4"

    wb.save(out_path)

    n_res = len(seeds.get("resources", []))
    n_int = len(seeds.get("integrations", []))
    print(f"Wrote {out_path}")
    print(f"  sheets: {len(wb.sheetnames)}  |  seeded RESOURCES: {n_res}  |  seeded INTEGRATIONS: {n_int}")


def main() -> None:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Rebuild the Visual Second Brain substrate workbook.")
    ap.add_argument("--spec", default=str(here / "substrate_spec.json"))
    ap.add_argument("--out", default=str(here.parent / "Visual_Second_Brain_Substrate.xlsx"))
    args = ap.parse_args()
    build(Path(args.spec), Path(args.out))


if __name__ == "__main__":
    main()
